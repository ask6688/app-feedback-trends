

"""Preserved legacy behavior."""

import csv
import datetime as dt
import hashlib
import io
import os
import re
from collections import Counter, OrderedDict, defaultdict

ALIAS_DATE = ["发表时间", "评论时间", "评论日期", "时间", "日期", "date", "comment_time", "comment_date", "created_at"]
ALIAS_STAR = ["评级", "星级", "评分", "分数", "星", "star", "rating", "score"]
ALIAS_TEXT = ["内容", "评价", "评论内容", "评论", "正文", "反馈内容", "comment", "content", "text", "review"]
ALIAS_TITLE = ["标题", "评论标题", "title", "subject"]
ALIAS_AUTHOR = ["作者", "评论人", "用户名", "用户", "昵称", "author", "user", "nickname", "user_name"]

ALIAS_CHANNEL = ["渠道", "来源渠道", "来源", "市场", "应用市场", "channel", "store", "source", "market"]
ALIAS_PLATFORM = ["平台", "操作系统", "系统", "platform", "device_type"]
ALIAS_MODEL = ["机型", "手机型号", "设备型号", "model", "device"]
ALIAS_VERSION = ["版本", "app版本", "应用版本", "version", "app_version"]
ALIAS_LIKES = ["点赞数", "点赞", "有用数", "likes", "like_count"]

ALIAS_CS_PERIOD = ["日期", "时间", "时间区间", "日期区间", "统计日期", "周期", "period", "date", "date_range"]
ALIAS_CS_CONV = ["原始对话", "对话内容", "会话内容", "聊天记录", "对话", "会话", "内容", "conversation", "content", "dialog"]
ALIAS_CS_TRIGGER = ["在线客服.触发用户数", "触发用户数", "触发用户", "trigger_users"]

AND_CHANNELS = ["OPPO", "VIVO", "华为", "小米", "魅族"]
AND_CHANNEL_ALIASES = {
    "oppo": "OPPO", "vivo": "VIVO", "huawei": "华为", "honor": "华为",
    "华为": "华为", "小米": "小米", "xiaomi": "小米", "redmi": "小米",
    "魅族": "魅族", "meizu": "魅族", "flyme": "魅族",
}
PLATFORM_ALIASES = {
    "ios": "IOS", "iphone": "IOS", "苹果": "IOS", "app store": "IOS", "appstore": "IOS",
    "and": "AND", "android": "AND", "安卓": "AND", "安卓市场": "AND",
}

CS_LABEL_HINTS = ("阶段汇总", "汇总", "合计", "总计", "小计")

_WS_RE = re.compile(r"\s+")
_WS_ANY_RE = re.compile(r"[\s\u200b\ufeff\u3000\xa0]+")
_PUNCT_MAP = str.maketrans({
    "，": ",", "。": ".", "！": "!", "？": "?", "：": ":", "；": ";",
    "（": "(", "）": ")", "【": "[", "】": "]", "「": '"', "」": '"',
    "『": '"', "』": '"', "、": ",", "～": "~", "－": "-", "—": "-", "…": ".",
})

def clean_text(v):
    """Preserved legacy behavior."""
    if v is None:
        return ""
    s = str(v)
    s = s.replace("\u200b", "").replace("\ufeff", "")
    s = _WS_RE.sub(" ", s)
    return s.strip()

def norm_ws(v):
    """强键用的文本归一：空白折叠 + strip（与历史数据口径一致）。"""
    return clean_text(v)

_PLAT_DEL_RE = re.compile(r"^\s*[（(]\s*该条评论已经被删除\s*[)）]\s*")
_PLAT_REPLY_RE = re.compile(r"\s*-{3,}\s*开发者回复[:：][\s\S]*$")

def strip_platform_markers(v):
    """剥掉平台对「内容」的改写：已删除前缀 + 开发者回复尾段。仅用于生成去重键。"""
    s = norm_ws(v)
    s = _PLAT_REPLY_RE.sub("", s)
    s = _PLAT_DEL_RE.sub("", s)
    return s.strip()

def key_text(v):
    """市场侧去重键的文本口径（强键与弱键共用）：先剥平台改写，再折叠空白。"""
    return strip_platform_markers(v)

def platform_form(v):
    """记录文本的形态标签，用于区分「真重复」与「源表本来就多行」。"""
    t = norm_ws(v)
    tags = []
    if _PLAT_DEL_RE.match(t):
        tags.append("已删除标记")
    if _PLAT_REPLY_RE.search(t):
        tags.append("含开发者回复")
    return "/".join(tags) if tags else "原始"

def _author_of(r):
    """记录的作者名（历史数据没有该字段时返回空串）。"""
    return norm_ws(r.get("author") or "")

def is_true_duplicate(a, b):
    """Preserved legacy behavior."""

    if key_text(a.get("comment_text")) != key_text(b.get("comment_text")):
        return False
    if norm_ws(a.get("comment_text")) == norm_ws(b.get("comment_text")):
        return False
    aa, ab = _author_of(a), _author_of(b)
    return bool(aa) and aa == ab

def scan_platform_duplicates(records):
    """自动识别「同一评论被平台改写文本后重复入库」。

    先按强键（platform|渠道|日期|星级|剥标记文本）分组，组内 >1 条时**按作者判定**：
      · true_dup    组内文本形态不同 + 存在作者明确相同的记录
                    → 真重复（同一条评论被算了两次），需处理
      · distinct    组内文本形态不同 + 作者都明确但互不相同
                    → 同文本不同用户，**不是重复，不要删**
      · unverified  组内文本形态不同 + 作者信息不足（历史数据未保留作者）
                    → 无法自动判定，只列出来人工核对；补齐作者后即可自动判定
      · same_text   组内文本形态全一致
                    → 源表本来就多行，不算重复

    返回 {"true_dup": [...], "distinct": [...], "unverified": [...], "same_text": [...],
          "counts": {...}}，每类已截断到 200 组。
    """
    groups = defaultdict(list)
    for r in records:
        text = r.get("comment_text")
        if not text:
            continue
        k = (r.get("platform") or "", r.get("and_channel") or "",
             r.get("comment_date") or "", str(r.get("star")), key_text(text))
        groups[k].append(r)

    true_dup, distinct, unverified, same_text = [], [], [], []
    for k, v in groups.items():
        if len(v) < 2:
            continue
        forms = {platform_form(x.get("comment_text")) for x in v}
        item = {
            "platform": k[0], "and_channel": k[1], "comment_date": k[2], "star": k[3],
            "text": k[4][:120],
            "records": [{"record_id": x.get("record_id"),
                         "form": platform_form(x.get("comment_text")),
                         "author": _author_of(x) or None,
                         "text_len": len(norm_ws(x.get("comment_text")))} for x in v],
        }
        if len(forms) == 1:
            same_text.append(item)
            continue

        by_author = defaultdict(list)
        unknown = 0
        for x in v:
            a = _author_of(x)
            if a:
                by_author[a].append(x)
            else:
                unknown += 1

        dup_authors = [a for a, xs in by_author.items()
                       if len({platform_form(x.get("comment_text")) for x in xs}) > 1]

        if dup_authors:
            item["reason"] = "作者相同（%s）且文本形态不同 → 同一条评论的两种形态" % "、".join(sorted(dup_authors))
            true_dup.append(item)
        elif unknown:
            item["reason"] = "作者信息不足（%d 条无作者），无法自动判定，请人工核对" % unknown
            unverified.append(item)
        elif len(by_author) > 1:
            item["reason"] = "作者不同（%s）→ 同文本不同用户，不是重复" % "、".join(sorted(by_author))
            distinct.append(item)
        else:
            item["reason"] = "无形态差异以外的区分依据"
            unverified.append(item)

    return {
        "true_dup": true_dup[:200],
        "distinct": distinct[:200],
        "unverified": unverified[:200],
        "same_text": same_text[:200],
        "counts": {
            "true_dup_groups": len(true_dup),
            "true_dup_extra_rows": sum(len(x["records"]) - 1 for x in true_dup),
            "distinct_groups": len(distinct),
            "distinct_extra_rows": sum(len(x["records"]) - 1 for x in distinct),
            "unverified_groups": len(unverified),
            "unverified_extra_rows": sum(len(x["records"]) - 1 for x in unverified),
            "same_text_groups": len(same_text),
            "same_text_extra_rows": sum(len(x["records"]) - 1 for x in same_text),
        },
    }

def norm_loose(v):
    """弱键用的文本归一：去掉所有空白，并统一全/半角标点。"""
    if v is None:
        return ""
    s = str(v).replace("\u200b", "").replace("\ufeff", "")
    s = _WS_ANY_RE.sub("", s)
    s = s.translate(_PUNCT_MAP)
    return s.strip()

DATE_RE = re.compile(r"(\d{4})\s*[-/.]\s*(\d{1,2})\s*[-/.]\s*(\d{1,2})")
TIME_RE = re.compile(r"(\d{1,2}):(\d{2})(?::(\d{2}))?")

CS_TAG_RE = re.compile(r"<br\s*/?>|<[a-zA-Z/!][^\n<>]{0,4000}>")
CS_TAG_TAIL_RE = re.compile(r"<[a-zA-Z/!][^\n<>]{0,4000}(?=\n|$)")
CS_ENTITY = [("&nbsp;", " "), ("&amp;", "&"), ("&lt;", "<"), ("&gt;", ">"),
             ("&quot;", '"'), ("&#39;", "'"), ("&apos;", "'")]

CS_WS_RE = re.compile(r"[ \t\u3000\xa0]+")
CS_USER_PREFIX = "用户："
CS_BOILER_RE = re.compile(
    r"已接入人工"
    r"|感谢您的咨询|祝您生活愉快|因咨询量较大|麻烦您留在咨询页面|请保持对话在线状态"
    r"|正在为您(验证|查询|核实|处理|排查)"
    r"|请您(提供|点击|进入|查看|确认|稍等|耐心)"
    r"|该情况是您|您的问题已升级|需转接至专员|转接至专员为您|专员会在\s*\d+\s*个自然日"
    r"|无法(定位|判断)到您的|请提供(以下|您)"
)

def cs_clean(t):
    """Preserved legacy behavior."""
    t = CS_TAG_RE.sub(" ", t)
    t = CS_TAG_TAIL_RE.sub(" ", t)
    for k, v in CS_ENTITY:
        t = t.replace(k, v)
    return CS_WS_RE.sub(" ", t).strip()

CS_CANON_RE = re.compile(r"[ \t\u3000\xa0?]+")
CS_SKEL_RE = re.compile(r"[^\u4e00-\u9fffA-Za-z0-9]+")

def cs_canon(t):
    """客服文本主键归一：丢弃 CR，把「空白类 + ASCII ?」折叠为单空格。"""
    return CS_CANON_RE.sub(" ", (t or "").replace("\r", "")).strip()

def cs_skel(t):
    """客服文本兜底归一：只保留中文与字母数字，忽略全部标点/符号/空白。"""
    return CS_SKEL_RE.sub("", (t or "").replace("\r", ""))

def detect_lossy_export(path):
    """判定这份客服导出是否为「字符集有损重导出」。

    这是 skeleton 级（L3）兜底匹配的**前置条件 (a)** —— 只有确认来源有损，
    才允许用骨架键把命中记录自动判为 existing_lossy；否则一律不启用。
    """
    name = os.path.basename(path).lower()
    if name.startswith("event_analysis"):
        return True, "文件名匹配 event_analysis*.csv（已确认该导出为有损重导出）"
    try:
        with open(path, "rb") as f:
            head = f.read(2_000_000)
    except Exception as e:
        return False, "无法读取文件内容：%s" % e
    s = head.decode("utf-8", errors="replace")
    crlf = s.count("\r\n")
    q = s.count("?")
    if crlf > 100 and q > 0:
        return True, ("内容特征：出现 %d 个 CRLF 与 %d 个 ASCII '?'"
                      "（有损重导出的典型特征）" % (crlf, q))
    return False, "未检测到有损重导出特征（CRLF %d / '?' %d）" % (crlf, q)

def norm_date(v):
    """Preserved legacy behavior."""
    if v is None:
        return None
    if hasattr(v, "strftime"):
        try:
            return v.strftime("%Y-%m-%d")
        except Exception:
            pass
    s = str(v).strip()
    if not s:
        return None
    m = DATE_RE.search(s)
    if not m:
        return None
    y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if not (2000 <= y <= 2100 and 1 <= mo <= 12 and 1 <= d <= 31):
        return None
    try:
        return dt.date(y, mo, d).isoformat()
    except ValueError:
        return None

def norm_time(v):
    """从原始时间串里抽出 HH:MM:SS（保留精度，便于人工核对/未来做更细的键）。"""
    if v is None:
        return None
    if hasattr(v, "strftime"):
        try:
            return v.strftime("%H:%M:%S")
        except Exception:
            return None
    s = str(v)
    m = TIME_RE.search(s)
    if not m:
        return None
    hh, mm = int(m.group(1)), int(m.group(2))
    ss = int(m.group(3)) if m.group(3) else 0
    if not (0 <= hh < 24 and 0 <= mm < 60 and 0 <= ss < 60):
        return None
    return "%02d:%02d:%02d" % (hh, mm, ss)

def norm_star(v):
    """Preserved legacy behavior."""
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    m = re.search(r"(\d+)", s)
    if not m:
        return None
    n = int(m.group(1))
    return n if 1 <= n <= 5 else None

def parse_period(v):
    """Preserved legacy behavior."""

    if v is None:
        return None, None, None, None
    if hasattr(v, "strftime"):
        s = v.strftime("%Y-%m-%d")
    else:
        s = str(v)
    raw = s.strip()
    if not raw:
        return None, None, None, None
    found = DATE_RE.findall(raw)
    if not found:
        return None, None, None, raw
    dates = []
    for y, mo, d in found:
        iso = norm_date("%s-%s-%s" % (y, mo, d))
        if iso:
            dates.append(iso)
    if not dates:
        return None, None, None, raw
    start = dates[0]
    end = dates[1] if len(dates) > 1 else dates[0]
    if end < start:
        start, end = end, start
    return start, end, "%s ～ %s" % (start, end), raw

def overlaps(a_start, a_end, b_start, b_end):
    """两个闭区间是否有重叠（任一端缺失 → False）。"""
    if not (a_start and a_end and b_start and b_end):
        return False
    return a_start <= b_end and b_start <= a_end

def short_hash(*parts):
    h = hashlib.sha1("\x1f".join("" if p is None else str(p) for p in parts).encode("utf-8"))
    return h.hexdigest()[:16]

def now_str():
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

def _norm_header(h):
    return _WS_ANY_RE.sub("", str(h or "")).lower()

def _find_col(header_norm, aliases, exclude=None):
    """在归一化表头里按别名精确匹配，再退化为包含匹配。返回列下标或 None。"""
    alias_norm = [_norm_header(a) for a in aliases]
    exclude_norm = [_norm_header(e) for e in (exclude or [])]
    for a in alias_norm:
        if a in header_norm:
            i = header_norm.index(a)
            if not any(e and e in header_norm[i] for e in exclude_norm):
                return i
    for i, h in enumerate(header_norm):
        if not h:
            continue
        if any(e and e in h for e in exclude_norm):
            continue
        for a in alias_norm:
            if a and (a in h or h in a) and len(a) >= 2:
                return i
    return None

def _read_xlsx(path):
    """Preserved legacy behavior."""

    try:
        import openpyxl
    except ImportError:
        raise SystemExit(
            "缺少 openpyxl，无法读取 .xlsx。\n"
            "请用带 openpyxl 的解释器运行，例如：\n"
            "  python -m pip install -r requirements.txt\n"
            "或先安装：pip install openpyxl"
        )
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    out = []
    try:
        for name in wb.sheetnames:
            ws = wb[name]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            out.append((name, rows))
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return out

def _read_csv(path):
    with open(path, "rb") as handle:
        raw = handle.read()
    text = None
    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("utf-8", errors="replace")

    text = text.replace("\r\r\n", "\n").replace("\r\n", "\n")
    rows = [r for r in csv.reader(io.StringIO(text), delimiter="\t" if str(path).lower().endswith(".tsv") else ",")]
    return [(os.path.basename(path), rows)]

def read_tables(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    if ext in (".csv", ".tsv"):
        return _read_csv(path)
    if ext == ".xls":
        raise SystemExit("不支持 .xls（老格式）：%s\n请另存为 .xlsx 或 .csv 后再导入。" % path)
    raise SystemExit("不支持的文件类型：%s" % path)

def find_header(rows, named_groups, max_scan=25):
    """在 rows 前 max_scan 行里找「同时命中多组别名」的那一行作为表头。

    named_groups: list[(语义名, [别名...], [排除别名...])]
    返回 (header_index, {语义名: 列下标})；找不到返回 (None, None)。
    """
    best = None
    for i, r in enumerate(rows[:max_scan]):
        header_norm = [_norm_header(c) for c in r]
        if not any(header_norm):
            continue
        hits = {}
        score = 0
        for grp in named_groups:
            name, aliases = grp[0], grp[1]
            exclude = grp[2] if len(grp) > 2 else None
            ci = _find_col(header_norm, aliases, exclude)
            if ci is not None:
                hits[name] = ci
                score += 1
        if score and (best is None or score > best[0]):
            best = (score, i, hits)
    if not best or best[0] < 1:
        return None, None
    return best[1], best[2]

def sniff_type(sheet_name, header_norm, filename):
    """粗判一个 sheet 属于哪类数据：'cs' / 'market' / None。"""
    blob = _norm_header(sheet_name) + "|" + _norm_header(filename) + "|" + "|".join(header_norm)
    has_conv = any(_norm_header(a) in header_norm for a in ALIAS_CS_CONV)
    has_star = any(_norm_header(a) in header_norm for a in ALIAS_STAR)
    if has_star:
        return "market"
    if has_conv and ("客服" in blob or "会话" in blob or "对话" in blob):
        return "cs"
    if has_conv:
        return "cs"
    return None

def detect_platform_channel(row, cols, sheet_name, filename, declared, channel_candidate_cols=()):
    """Preserved legacy behavior."""

    platform = None
    channel = ""

    if "platform" in cols:
        v = row[cols["platform"]] if cols["platform"] < len(row) else None
        p = PLATFORM_ALIASES.get(_norm_header(v))
        if p:
            platform = p

    if "channel" in cols:
        v = row[cols["channel"]] if cols["channel"] < len(row) else None
        c = AND_CHANNEL_ALIASES.get(_norm_header(v))
        if c:
            channel = c
            platform = platform or "AND"
        else:
            p = PLATFORM_ALIASES.get(_norm_header(v))
            if p:
                platform = platform or p

    if not channel:
        blob = _norm_header(sheet_name) + "|" + _norm_header(filename)
        for key, c in AND_CHANNEL_ALIASES.items():
            if key and key in blob:
                channel = c
                platform = platform or "AND"
                break

    if not channel and declared == "and":
        for ci in channel_candidate_cols:
            if ci is None or ci >= len(row):
                continue
            c = AND_CHANNEL_ALIASES.get(_norm_header(row[ci]))
            if c:
                channel = c
                platform = platform or "AND"
                break

    if platform is None:
        if declared == "ios":
            platform = "IOS"
        elif declared == "and":
            platform = "AND"
    if platform == "IOS":
        channel = ""
    return platform, channel

class MarketIndex(object):
    """市场侧历史索引：强键 / 弱键多重集 + 冻结副本（判断键是否存在于历史）。"""

    def __init__(self, mode="multiset"):
        self.mode = mode
        self.strong = Counter()
        self.strong_all = Counter()
        self.weak = Counter()
        self.weak_examples = {}
        self.total = 0

    def add(self, strong, weak, record):
        self.strong[strong] += 1
        self.strong_all[strong] += 1
        self.weak[weak] += 1
        self.weak_examples.setdefault(weak, []).append(record)
        self.total += 1

    def take(self, strong, weak):
        if self.strong.get(strong, 0) > 0:
            if self.mode == "multiset":
                self.strong[strong] -= 1
                if self.weak.get(weak, 0) > 0:
                    self.weak[weak] -= 1
            return "existing"
        if self.weak.get(weak, 0) > 0:
            if self.mode == "multiset":
                self.weak[weak] -= 1
            return "updated"
        return "new"

    def pick_weak_example(self, weak):
        lst = self.weak_examples.get(weak)
        return lst[0] if lst else None

class CsIndex(object):
    """Preserved legacy behavior."""

    def __init__(self, mode="multiset", time_mode="overlap"):
        self.mode = mode
        self.time_mode = time_mode
        self.by_text = defaultdict(list)
        self.by_canon = defaultdict(list)
        self.by_skel = defaultdict(list)
        self.by_loose = defaultdict(list)
        self.strong_all = Counter()
        self.text_all = Counter()
        self.canon_all = Counter()
        self.used = set()
        self.total = 0

    def _win_ok(self, i_start, i_end, p_start, p_end):
        """时间窗是否算命中。ignore 模式下不做时间过滤。"""
        if self.time_mode == "ignore":
            return True
        return overlaps(i_start, i_end, p_start, p_end)

    def add(self, rec):
        i = self.total
        rec = dict(rec)
        rec["__idx"] = i
        clean = rec.get("cs_text_clean") or ""
        self.by_text[clean].append(rec)
        self.by_canon[cs_canon(clean)].append(rec)
        self.by_skel[cs_skel(clean)].append(rec)
        self.by_loose[norm_loose(rec.get("cs_text"))].append(rec)
        self.strong_all[(rec.get("period_start") or "", rec.get("period_end") or "",
                         clean)] += 1
        self.text_all[clean] += 1
        self.canon_all[cs_canon(clean)] += 1
        self.total += 1

    def _candidates(self, pool, i_start, i_end, consume=True):
        """返回池子里「未被核销 且 时间窗可命中」的记录；consume=True 时核销第一条。"""
        out = []
        for rec in pool:
            if rec["__idx"] in self.used:
                continue
            if self._win_ok(i_start, i_end, rec.get("period_start"), rec.get("period_end")):
                out.append(rec)
        if consume and out:
            self.used.add(out[0]["__idx"])
        return out

    def _pick(self, pool, i_start, i_end):
        got = self._candidates(pool, i_start, i_end, consume=True)
        return got[0] if got else None

    def take(self, i_start, i_end, text_clean, text_loose,
             text_canon=None, text_skel=None, allow_skeleton=False):
        """返回 (verdict, payload, level)。

        verdict ∈ existing | existing_lossy | ambiguous_lossy | updated | new
        payload = 命中的历史记录（ambiguous_lossy 时为候选列表）
        level   ∈ strict | canon | skeleton | loose | None
        """
        if text_canon is None:
            text_canon = cs_canon(text_clean)
        if text_skel is None:
            text_skel = cs_skel(text_clean)

        if self.mode == "set":
            for pool, verdict, level in ((self.by_text.get(text_clean, []), "existing", "strict"),
                                         (self.by_canon.get(text_canon, []), "existing", "canon"),
                                         (self.by_loose.get(text_loose, []), "updated", "loose")):
                for rec in pool:
                    if self._win_ok(i_start, i_end, rec.get("period_start"), rec.get("period_end")):
                        return verdict, rec, level
            if allow_skeleton:
                got = self._candidates(self.by_skel.get(text_skel, []), i_start, i_end, consume=False)
                if got:
                    if len({(r.get("cs_text_clean") or "") for r in got}) > 1:
                        return "ambiguous_lossy", got, "skeleton"
                    return "existing_lossy", got[0], "skeleton"
            return "new", None, None

        rec = self._pick(self.by_text.get(text_clean, []), i_start, i_end)
        if rec is not None:
            return "existing", rec, "strict"

        rec = self._pick(self.by_canon.get(text_canon, []), i_start, i_end)
        if rec is not None:
            return "existing", rec, "canon"

        if allow_skeleton:
            got = self._candidates(self.by_skel.get(text_skel, []), i_start, i_end, consume=False)
            if got:
                distinct = {(r.get("cs_text_clean") or "") for r in got}
                if len(distinct) > 1:

                    return "ambiguous_lossy", got, "skeleton"
                self.used.add(got[0]["__idx"])
                return "existing_lossy", got[0], "skeleton"

        rec = self._pick(self.by_loose.get(text_loose, []), i_start, i_end)
        if rec is not None:
            return "updated", rec, "loose"
        return "new", None, None

    def capacity(self, i_start, i_end, text_clean):
        """历史里「同文本（+时间约束）」的记录总条数（不扣除已核销的）。"""
        k = cs_canon(text_clean)
        return sum(1 for rec in self.by_canon.get(k, [])
                   if self._win_ok(i_start, i_end, rec.get("period_start"), rec.get("period_end")))

    def weak_example(self, text_loose, i_start, i_end):
        for rec in self.by_loose.get(text_loose, []):
            if self._win_ok(i_start, i_end, rec.get("period_start"), rec.get("period_end")):
                return rec
        return None

def build_market_records(tables, declared, source_path, warnings, sheet_stats):
    """把 IOS / AND 的表格转成统一记录。返回 list[dict]。"""
    records = []
    fname = os.path.basename(source_path)
    for sheet_name, rows in tables:
        if not rows:
            continue
        hdr_idx, cols = find_header(rows, [
            ("date", ALIAS_DATE), ("star", ALIAS_STAR), ("text", ALIAS_TEXT),
        ])
        if hdr_idx is None or "text" not in cols:
            warnings.append("跳过 sheet「%s」：未识别到表头（需要 时间/星级/内容 类列）" % sheet_name)
            continue
        header_raw = [clean_text(c) for c in rows[hdr_idx]]
        header_norm = [_norm_header(c) for c in rows[hdr_idx]]
        stype = sniff_type(sheet_name, header_norm, source_path)
        if stype == "cs":
            warnings.append("sheet「%s」看起来是客服数据，但被当作市场侧导入 —— 请确认" % sheet_name)
        for name, aliases in (("title", ALIAS_TITLE), ("author", ALIAS_AUTHOR),
                              ("model", ALIAS_MODEL), ("version", ALIAS_VERSION),
                              ("likes", ALIAS_LIKES), ("platform", ALIAS_PLATFORM)):
            if name in cols:
                continue
            ci = _find_col(header_norm, aliases)
            if ci is not None:
                cols[name] = ci
        if "channel" not in cols:
            ci = _find_col(header_norm, ALIAS_CHANNEL)
            if ci is not None:
                cols["channel"] = ci

        cand = []
        for nm in ("channel", "model", "platform"):
            ci = cols.get(nm)
            if ci is not None and ci not in cand:
                cand.append(ci)

        n_ok = 0
        n_bad_date = 0
        n_empty = 0
        dates = []
        for ri, r in enumerate(rows[hdr_idx + 1:], start=hdr_idx + 2):
            def g(key):
                ci = cols.get(key)
                return r[ci] if (ci is not None and ci < len(r)) else None

            text = norm_ws(g("text"))
            title = norm_ws(g("title"))
            raw_date = g("date")
            d = norm_date(raw_date)
            st = norm_star(g("star"))
            if d is None and st is None and not text and not title:
                n_empty += 1
                continue
            if d is None:
                n_bad_date += 1
            platform, channel = detect_platform_channel(r, cols, sheet_name, source_path,
                                                        declared, cand)
            raw = OrderedDict()
            for ci2, hname in enumerate(header_raw):
                if hname and ci2 < len(r) and r[ci2] is not None:
                    raw[hname] = clean_text(r[ci2])
            rec = OrderedDict()
            rec["data_domain"] = "market"
            rec["source_type"] = declared
            rec["platform"] = platform
            rec["and_channel"] = channel
            rec["comment_date"] = d
            rec["comment_time"] = norm_time(raw_date)
            rec["comment_time_raw"] = clean_text(raw_date)
            rec["star"] = st
            rec["title"] = title
            rec["comment_text"] = text
            rec["author"] = norm_ws(g("author"))
            rec["model"] = norm_ws(g("model"))
            rec["app_version"] = norm_ws(g("version"))
            rec["likes"] = clean_text(g("likes"))
            rec["is_negative"] = bool(st is not None and st <= 3)
            rec["source_file"] = fname
            rec["source_path"] = source_path
            rec["source_sheet"] = sheet_name
            rec["source_row"] = ri
            rec["raw"] = raw
            records.append(rec)
            n_ok += 1
            if d:
                dates.append(d)
        sheet_stats.append(OrderedDict([
            ("source_file", fname), ("source_sheet", sheet_name), ("data_type", "market"),
            ("rows_parsed", n_ok), ("rows_empty_skipped", n_empty),
            ("date_min", min(dates) if dates else None),
            ("date_max", max(dates) if dates else None),
        ]))
        if n_bad_date:
            warnings.append("sheet「%s」有 %d 行日期解析失败（已保留，comment_date=null）" % (sheet_name, n_bad_date))
    return records

def build_cs_records(tables, source_path, warnings, sheet_stats, anomalies):
    """把客服表格转成统一记录。非日期行（如「阶段汇总」）归入 anomalies。"""
    records = []
    fname = os.path.basename(source_path)
    lossy, lossy_reason = detect_lossy_export(source_path)
    if lossy:
        warnings.append("客服源「%s」判定为**字符集有损重导出** → 启用 skeleton 级兜底匹配：%s"
                        % (fname, lossy_reason))
    else:
        warnings.append("客服源「%s」未检出字符集有损 → **不启用** skeleton 级兜底匹配：%s"
                        % (fname, lossy_reason))
    for sheet_name, rows in tables:
        if not rows:
            continue
        hdr_idx, cols = find_header(rows, [("period", ALIAS_CS_PERIOD), ("conv", ALIAS_CS_CONV)])
        if hdr_idx is None or "conv" not in cols:
            warnings.append("跳过 sheet「%s」：未识别到客服表头（需要 时间/对话内容 类列）" % sheet_name)
            continue
        header_raw = [clean_text(c) for c in rows[hdr_idx]]
        header_norm = [_norm_header(c) for c in rows[hdr_idx]]
        ci = _find_col(header_norm, ALIAS_CS_TRIGGER)
        if ci is not None and "trigger" not in cols:
            cols["trigger"] = ci
        stype = sniff_type(sheet_name, header_norm, source_path)
        if stype == "market":
            warnings.append("sheet「%s」看起来是市场侧数据（有星级列），却被当作客服导入 —— 请确认" % sheet_name)

        n_ok = 0
        n_empty = 0
        n_label = 0
        n_rollup = 0
        label_counter = Counter()
        dates = []
        body = rows[hdr_idx + 1:]

        ci_period = cols.get("period")
        ci_conv = cols.get("conv")
        dated_convs = set()
        for r0 in body:
            cv0 = r0[ci_conv] if (ci_conv is not None and ci_conv < len(r0)) else None
            if cv0 is None or not str(cv0).strip():
                continue
            pv0 = r0[ci_period] if (ci_period is not None and ci_period < len(r0)) else None
            if parse_period(pv0)[0] is not None:
                dated_convs.add(str(cv0))
        for ri, r in enumerate(body, start=hdr_idx + 2):
            def g(key):
                ci2 = cols.get(key)
                return r[ci2] if (ci2 is not None and ci2 < len(r)) else None

            conv = g("conv")
            conv_s = "" if conv is None else str(conv)
            if not conv_s.strip():
                n_empty += 1
                continue
            p_start, p_end, p_label, p_raw = parse_period(g("period"))
            raw = OrderedDict()
            for ci2, hname in enumerate(header_raw):
                if hname and ci2 < len(r) and r[ci2] is not None:
                    raw[hname] = clean_text(r[ci2])

            raw_users = [ln.strip()[len(CS_USER_PREFIX):].strip()
                         for ln in conv_s.split("\n")
                         if ln.strip().startswith(CS_USER_PREFIX)]
            users = [u for u in raw_users if not CS_BOILER_RE.search(u)]
            if users:
                raw_user = "\n".join(users).strip()
            elif raw_users:
                raw_user = ""
            else:
                raw_user = conv_s.strip()
                warnings.append("sheet「%s」第 %d 行没有「用户：」前缀，已按原文整段处理" % (sheet_name, ri))

            if p_start is None:
                if conv_s.strip() and conv_s in dated_convs:

                    n_rollup += 1
                    continue

                n_label += 1
                label_counter[p_raw or "(空)"] += 1
                an = OrderedDict()
                an["data_domain"] = "cs"
                an["source_type"] = "cs"
                an["anomaly"] = "unparsable_period"
                an["period_raw"] = p_raw
                an["source_file"] = fname
                an["source_path"] = source_path
                an["source_sheet"] = sheet_name
                an["source_row"] = ri
                an["cs_text"] = raw_user
                an["cs_text_clean"] = cs_clean(raw_user)
                an["raw"] = raw
                anomalies.append(an)
                continue

            rec = OrderedDict()
            rec["data_domain"] = "cs"
            rec["source_type"] = "cs"
            rec["period"] = p_label
            rec["period_start"] = p_start
            rec["period_end"] = p_end
            rec["period_raw"] = p_raw
            rec["cs_text"] = raw_user
            rec["cs_text_clean"] = cs_clean(raw_user)
            rec["user_turns"] = len(users)
            rec["user_turns_raw"] = len(raw_users)
            rec["trigger_users"] = clean_text(g("trigger"))
            rec["source_file"] = fname
            rec["source_path"] = source_path
            rec["source_sheet"] = sheet_name
            rec["source_row"] = ri
            rec["source_is_lossy_export"] = lossy
            rec["source_lossy_reason"] = lossy_reason
            rec["raw"] = raw
            records.append(rec)
            n_ok += 1
            dates.extend([p_start, p_end])
        sheet_stats.append(OrderedDict([
            ("source_file", fname), ("source_sheet", sheet_name), ("data_type", "cs"),
            ("rows_parsed", n_ok), ("rows_empty_skipped", n_empty),
            ("rows_rollup_skipped", n_rollup),
            ("rows_unparsable_period", n_label),
            ("date_min", min(dates) if dates else None),
            ("date_max", max(dates) if dates else None),
        ]))
        if n_rollup:
            warnings.append("sheet「%s」有 %d 行是「阶段汇总」等汇总段落的**重复罗列**"
                            "（同一会话已在日期段落出现过），已直接丢弃、不计入异常"
                            % (sheet_name, n_rollup))
        if n_label:
            warnings.append("sheet「%s」有 %d 行「时间」列不是日期（%s），已归入 anomalies 不参与比对"
                            % (sheet_name, n_label,
                               ", ".join("%s×%d" % (k, v) for k, v in label_counter.most_common(5))))
    return records

def market_keys(rec):

    text = key_text(rec.get("comment_text"))
    strong = (rec.get("platform") or "", rec.get("and_channel") or "",
              rec.get("comment_date") or "", str(rec.get("star")), text)
    weak = (rec.get("platform") or "", rec.get("and_channel") or "",
            rec.get("comment_date") or "", text)
    return strong, weak

def key_str(key):
    return "|".join("" if p is None else str(p) for p in key)

def cs_key_str(rec, time_mode="overlap"):
    base = "canon:%s" % short_hash(cs_canon(rec.get("cs_text_clean") or ""))
    if time_mode == "ignore":
        return base
    return "%s~%s|%s" % (rec.get("period_start") or "", rec.get("period_end") or "", base)



def resolve_cs_time_mode(records, requested):
    """Preserved legacy behavior."""

    if requested != "auto":
        return requested, "命令行显式指定 --cs-time-mode=%s" % requested
    labels = set()
    for r in records:
        if r.get("period_start"):
            labels.add((r.get("period_start"), r.get("period_end")))
    if len(labels) == 1:
        s, e = next(iter(labels))
        return "ignore", ("输入只有一个日期取值（%s%s），判定为「批次截止日 / 阶段标签」"
                          "而非会话发生日 → 不按时间过滤，只按文本匹配"
                          % (s, "" if s == e else " ~ " + str(e)))
    if not labels:
        return "ignore", "输入没有可解析的日期 → 不按时间过滤"
    return "overlap", "输入有 %d 个不同日期/区间 → 按时间窗重叠匹配" % len(labels)

def diff_market(hist, rec):
    strong, weak = market_keys(rec)
    verdict = hist.take(strong, weak)
    if verdict == "updated":
        old = hist.pick_weak_example(weak)
        changed = []
        if old is not None:
            if str(old.get("star")) != str(rec.get("star")):
                changed.append({"field": "star", "old": old.get("star"), "new": rec.get("star")})
            if norm_ws(old.get("comment_text")) != norm_ws(rec.get("comment_text")):
                changed.append({"field": "comment_text", "old": norm_ws(old.get("comment_text")),
                                "new": norm_ws(rec.get("comment_text"))})
        return verdict, {"matched_record_id": (old or {}).get("record_id"), "changed_fields": changed}, old
    return verdict, None, None

def diff_cs(hist, rec):
    i_start, i_end = rec.get("period_start"), rec.get("period_end")
    text_clean = rec.get("cs_text_clean") or ""
    text_canon = cs_canon(text_clean)
    text_skel = cs_skel(text_clean)
    text_loose = norm_loose(rec.get("cs_text"))

    allow_skeleton = bool(rec.get("source_is_lossy_export"))
    verdict, old, level = hist.take(i_start, i_end, text_clean, text_loose,
                                    text_canon, text_skel, allow_skeleton=allow_skeleton)
    if verdict == "ambiguous_lossy":
        cands = old or []
        return verdict, {
            "match_level": "skeleton",
            "reason": "skeleton_ambiguous",
            "note": "骨架键命中了 %d 条**内容各不相同**的历史记录，无法唯一定位 → 不自动判已存在"
                    % len(cands),
            "candidate_count": len(cands),
            "candidates": [{
                "record_id": c.get("record_id"),
                "period": c.get("period"),
                "cs_text_clean": (c.get("cs_text_clean") or "")[:200],
            } for c in cands[:5]],
            "new_cs_text_clean": text_clean[:200],
        }, None
    if verdict == "existing_lossy":
        return verdict, {"match_level": "skeleton",
                         "matched_record_id": (old or {}).get("record_id"),
                         "reason": "lossy_charset_rewrite",
                         "note": "骨架键（中文+字母数字）完全一致且可唯一定位，"
                                 "仅空白/特殊字符被导出改写成 '?'",
                         "old_cs_text_clean": (old or {}).get("cs_text_clean", "")[:200],
                         "new_cs_text_clean": text_clean[:200]}, old
    if verdict == "updated":
        if old is None:
            old = hist.weak_example(text_loose, i_start, i_end)
        changed = []
        if old is not None:
            if (old.get("cs_text_clean") or "") != text_clean:
                changed.append({"field": "cs_text_clean",
                                "old": (old.get("cs_text_clean") or "")[:200],
                                "new": text_clean[:200]})
            if str(old.get("user_turns")) != str(rec.get("user_turns")):
                changed.append({"field": "user_turns", "old": old.get("user_turns"),
                                "new": rec.get("user_turns")})

            if hist.time_mode != "ignore" and not overlaps(i_start, i_end, old.get("period_start"),
                                                           old.get("period_end")):
                changed.append({"field": "period",
                                "old": "%s ~ %s" % (old.get("period_start"), old.get("period_end")),
                                "new": "%s ~ %s" % (i_start, i_end)})
        return verdict, {"match_level": level, "matched_record_id": (old or {}).get("record_id"),
                         "changed_fields": changed}, old
    return verdict, ({"match_level": level} if level else None), old

DISTINGUISH_FIELDS = ["comment_time", "author", "model", "app_version", "likes",
                      "and_channel", "platform", "star", "comment_date"]

def analyze_input_collisions(records, kind):
    """统计输入文件内部的强键碰撞，并检查能否用其他原始字段区分。"""
    if kind == "cs":
        def keyfn(r):
            return (r.get("period_start"), r.get("period_end"), r.get("cs_text_clean") or "")
        fields = ["period_start", "period_end", "user_turns", "trigger_users"]
    else:
        def keyfn(r):
            return market_keys(r)[0]
        fields = DISTINGUISH_FIELDS

    groups = defaultdict(list)
    for r in records:
        groups[keyfn(r)].append(r)
    collided = {k: v for k, v in groups.items() if len(v) > 1}

    details = []
    n_distinguishable = 0
    for k, rows in sorted(collided.items(), key=lambda x: -len(x[1]))[:200]:
        sigs = set()
        for r in rows:
            sigs.add(tuple(str(r.get(f)) for f in fields))
        distinguishable = len(sigs) > 1
        if distinguishable:
            n_distinguishable += 1
        details.append(OrderedDict([
            ("count", len(rows)),
            ("key", key_str(k) if kind != "cs" else "%s~%s|%s" % (k[0], k[1], short_hash(k[2]))),
            ("distinguishable_by_other_fields", distinguishable),
            ("distinguishing_fields_available", [f for f in fields if len({str(r.get(f)) for r in rows}) > 1]),
            ("sample", [OrderedDict((f, r.get(f)) for f in fields) for r in rows[:3]]),
        ]))
    return OrderedDict([
        ("collision_groups", len(collided)),
        ("collision_redundant_rows", sum(len(v) - 1 for v in collided.values())),
        ("groups_distinguishable_by_other_fields", n_distinguishable),
        ("groups_not_distinguishable", len(collided) - n_distinguishable),
        ("details", details),
    ])

TYPE_LABEL = {"ios": "IOS", "and": "AND", "cs": "客服"}

def process(name, paths, hist, kind, preview_prefix, imported_at, cs_time_mode="auto"):
    """处理一类输入（可多文件），返回 (bucket, stats, collisions, sheet_stats, anomalies)。"""
    warnings = []
    sheet_stats = []
    anomalies = []
    records = []
    for p in paths:
        tables = read_tables(p)
        if kind == "cs":
            records.extend(build_cs_records(tables, p, warnings, sheet_stats, anomalies))
        else:
            records.extend(build_market_records(tables, name, p, warnings, sheet_stats))

    resolved_mode = None
    mode_reason = None
    if kind == "cs":
        resolved_mode, mode_reason = resolve_cs_time_mode(records, cs_time_mode)
        hist.time_mode = resolved_mode
        if resolved_mode == "ignore":
            warnings.append("客服「时间」列按**批次截止日**处理（不参与匹配）：%s" % mode_reason)
        else:
            warnings.append("客服「时间」列按**真实日期区间**处理：%s" % mode_reason)

    collisions = analyze_input_collisions(records, kind)

    n_intra_batch = 0
    if kind != "cs":
        seen = defaultdict(list)
        kept = []
        for rec in records:
            k = market_keys(rec)[0]
            if any(is_true_duplicate(prev, rec) for prev in seen[k]):
                n_intra_batch += 1
                continue
            seen[k].append(rec)
            kept.append(rec)
        records = kept
        if n_intra_batch:
            warnings.append(
                "批内重叠窗口：本次输入里有 %d 条与同批前面某条属于「同一评论的两种文本形态」"
                "（作者相同），已合并按 1 条计入；建议后续只喂**不重叠**的周窗口。" % n_intra_batch)

    bucket = {"existing": [], "existing_lossy": [], "ambiguous_lossy": [],
              "updated": [], "new": []}
    seq = 0
    n_dup_exhausted = 0
    for rec in records:
        if kind == "cs":
            verdict, detail, old = diff_cs(hist, rec)
            strong_desc = cs_key_str(rec, hist.time_mode)
            skel_desc = "skel:%s" % short_hash(cs_skel(rec.get("cs_text_clean") or ""))
            weak_desc = ("loose:%s" % short_hash(norm_loose(rec.get("cs_text")))) if hist.time_mode == "ignore" \
                else "%s~%s|%s" % (rec.get("period_start") or "", rec.get("period_end") or "",
                                   short_hash(norm_loose(rec.get("cs_text"))))
        else:
            verdict, detail, old = diff_market(hist, rec)
            strong, weak = market_keys(rec)
            strong_desc = key_str(strong)
            weak_desc = key_str(weak)
            skel_desc = None

        rec = OrderedDict(rec)
        if kind == "cs":
            rec["period_is_cutoff"] = (hist.time_mode == "ignore")
        rec["record_key"] = strong_desc
        rec["record_key_hash"] = short_hash(strong_desc)
        if skel_desc is not None:
            rec["skeleton_key"] = skel_desc
            rec["skeleton_key_hash"] = short_hash(skel_desc)
        rec["weak_key"] = weak_desc
        rec["weak_key_hash"] = short_hash(weak_desc)
        rec["dedup_status"] = verdict
        rec["match_level"] = (detail or {}).get("match_level")
        rec["matched_record_id"] = (old or {}).get("record_id")
        rec["imported_at"] = imported_at

        if verdict == "new":
            seq += 1
            rec["preview_id"] = "%s-NEW-%06d" % (preview_prefix, seq)
            exists = False
            if kind == "cs":

                exists = hist.capacity(rec.get("period_start"), rec.get("period_end"),
                                       rec.get("cs_text_clean") or "") > 0
            else:
                exists = hist.strong_all.get(market_keys(rec)[0], 0) > 0
            if exists:
                rec["new_reason"] = "duplicate_key_exhausted"
                rec["new_reason_note"] = ("该强键在历史中存在，但历史条数已被本次输入的其他行核销完；"
                                          "按多重集语义算新增")
                n_dup_exhausted += 1
            else:
                rec["new_reason"] = "not_in_history"
        if verdict == "existing_lossy":
            rec["match_note"] = "lossy_charset_rewrite"
        if verdict == "ambiguous_lossy":
            rec["match_note"] = "skeleton_ambiguous_needs_review"
            rec["ambiguous_candidates"] = (detail or {}).get("candidates")
            rec["ambiguous_candidate_count"] = (detail or {}).get("candidate_count")
        if verdict == "updated" and detail:
            rec["update_detail"] = detail
        bucket[verdict].append(rec)

    dates = []
    for r in bucket["new"]:
        if kind == "cs":
            dates.extend([d for d in (r.get("period_start"), r.get("period_end")) if d])
        else:
            if r.get("comment_date"):
                dates.append(r["comment_date"])

    stats = OrderedDict()
    stats["files"] = paths
    stats["input_total"] = len(records)
    stats["existing"] = len(bucket["existing"])
    stats["existing_lossy"] = len(bucket["existing_lossy"])
    stats["existing_total"] = len(bucket["existing"]) + len(bucket["existing_lossy"])
    stats["ambiguous_lossy"] = len(bucket["ambiguous_lossy"])
    stats["new"] = len(bucket["new"])
    stats["updated"] = len(bucket["updated"])
    stats["anomaly"] = len(anomalies)
    stats["match_levels"] = dict(Counter(r.get("match_level") for r in records
                                         if r.get("match_level")))
    stats["key_collision_groups_in_input"] = collisions["collision_groups"]
    stats["key_collision_rows_in_input"] = collisions["collision_redundant_rows"]
    stats["new_not_in_history"] = len(bucket["new"]) - n_dup_exhausted
    stats["new_duplicate_key_exhausted"] = n_dup_exhausted
    stats["new_date_min"] = min(dates) if dates else None
    stats["new_date_max"] = max(dates) if dates else None
    if kind == "cs":
        stats["cs_time_mode_resolved"] = resolved_mode
        stats["cs_time_mode_reason"] = mode_reason
    _exist = bucket["existing"] + bucket["existing_lossy"]
    stats["existing_date_min"] = min([r.get("comment_date") or r.get("period_start")
                                      for r in _exist if (r.get("comment_date") or r.get("period_start"))] or [None])
    stats["existing_date_max"] = max([r.get("comment_date") or r.get("period_end")
                                      for r in _exist if (r.get("comment_date") or r.get("period_end"))] or [None])
    stats["warnings"] = warnings
    return bucket, stats, collisions, sheet_stats, anomalies
